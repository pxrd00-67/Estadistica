#!/usr/bin/env python3
"""Ejecuta cada caso de prueba contra la aplicación real y llena docs/Pruebas_Calculadora.xlsx.

Uso (desde la raíz):  mvn -q compile && python3 tools/generar_excel.py
"""
import re, subprocess, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

OPS = {"suma": "a + b", "resta": "a - b", "multiplicacion": "a * b", "division": "a / b"}
DIV0 = "a / b -> Error: no se puede dividir entre cero."

# (id, función, descripción, entrada a, entrada b, resultado esperado)
CASOS = [
 ("TC-01", "suma", "Dos positivos", "5", "3", "8"),
 ("TC-02", "suma", "Dos negativos", "-5", "-3", "-8"),
 ("TC-03", "suma", "Sumar cero", "5", "0", "5"),
 ("TC-04", "suma", "Límite superior de int", "2147483647", "2147483647", "4294967294"),
 ("TC-05", "resta", "Dos positivos", "5", "3", "2"),
 ("TC-06", "resta", "Resultado negativo", "3", "5", "-2"),
 ("TC-07", "resta", "Dos negativos", "-5", "-3", "-2"),
 ("TC-08", "multiplicacion", "Dos positivos", "5", "3", "15"),
 ("TC-09", "multiplicacion", "Multiplicar por cero", "5", "0", "0"),
 ("TC-10", "multiplicacion", "Signos distintos", "-5", "3", "-15"),
 ("TC-11", "multiplicacion", "Límite superior de int", "2147483647", "2147483647", "4611686014132420609"),
 ("TC-12", "division", "División exacta", "6", "3", "2.0"),
 ("TC-13", "division", "División con decimales", "5", "2", "2.5"),
 ("TC-14", "division", "Cero entre número", "0", "7", "0.0"),
 ("TC-15", "division", "División negativa", "-5", "2", "-2.5"),
 ("TC-16", "division", "División entre cero", "5", "0", DIV0),
 ("TC-17", "division", "Cero entre cero", "0", "0", DIV0),
 ("TC-18", "entrada", "Entero válido", "42", "1", "OK"),
 ("TC-19", "entrada", "Texto no numérico se rechaza", "abc", "1", "ERROR"),
 ("TC-20", "entrada", "Decimal se rechaza", "3.5", "1", "ERROR"),
 ("TC-21", "entrada", "Fuera del rango de int se rechaza", "99999999999", "1", "ERROR"),
 ("TC-22", "entrada", "Entrada vacía", "", "1", "ERROR"),
]
try:
    from casos_extra import CASOS_EXTRA, OPS_EXTRA
    OPS.update(OPS_EXTRA); CASOS += CASOS_EXTRA
except ImportError:
    pass

def ejecutar(a, b):
    p = subprocess.run(["java", "-Dstdout.encoding=UTF-8", "-cp", "target/classes",
                        "co.estadistica.calculadora.Main"], input=f"{a}\n{b}\n",
                       capture_output=True, text=True, encoding="utf-8")
    return p.stdout

def obtenido(fn, a, b):
    out = ejecutar(a, b)
    if fn == "entrada":
        return ("OK" if "Error:" not in out else "ERROR"), None
    etiqueta = OPS[fn]
    for linea in out.splitlines():
        if etiqueta in linea:
            resto = linea[linea.index(etiqueta):]
            return (resto.split("=", 1)[1].strip() if " = " in resto else resto), None
    return "(sin salida)", None

wb = Workbook(); ws = wb.active; ws.title = "Pruebas"
enc = ["ID", "Función", "Descripción", "Entrada a", "Entrada b", "Resultado esperado",
       "Resultado obtenido", "Estado", "Fecha"]
ws.append(enc)
import datetime
hoy = datetime.date.today().isoformat()
fallos = 0
for id_, fn, desc, a, b, esp in CASOS:
    res = obtenido(fn, a, b)[0]
    ok = res == esp
    fallos += not ok
    ws.append([id_, fn, desc, a, b, esp, res, "PASA" if ok else "FALLA", hoy])
verde = PatternFill("solid", fgColor="C6EFCE"); rojo = PatternFill("solid", fgColor="FFC7CE")
borde = Border(*(Side(style="thin"),) * 4)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    c.alignment = Alignment(horizontal="center")
for fila in ws.iter_rows(min_row=2):
    for c in fila:
        c.border = borde; c.number_format = "@"
    fila[7].fill = verde if fila[7].value == "PASA" else rojo
for col, w in zip("ABCDEFGHI", (8, 16, 36, 14, 14, 44, 44, 9, 12)):
    ws.column_dimensions[col].width = w
n = len(CASOS)
ws.append([]); ws.append(["", "", "Total casos", n]); ws.append(["", "", "Pasan", n - fallos])
ws.append(["", "", "Fallan", fallos])
wb.save("docs/Pruebas_Calculadora.xlsx")
print(f"{n} casos, {n - fallos} pasan, {fallos} fallan")
sys.exit(1 if fallos else 0)
